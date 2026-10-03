"""JARVIS File Intelligence.

This module is deliberately additive: it provides controlled file context and
document operations while reusing the existing registry and AI service.
Uploaded files are stored below ``data/file_intelligence`` and never in the
source tree.
"""

from __future__ import annotations

import csv
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from core.paths import DATA
from core.registry import register


ROOT = DATA / "file_intelligence"
FILES_DIR = ROOT / "files"
INDEX_PATH = ROOT / "index.json"
MAX_FILE_BYTES = 500 * 1024 * 1024
MAX_TEXT_BYTES = 12 * 1024 * 1024

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".py", ".js", ".ts", ".jsx", ".tsx",
    ".html", ".htm", ".css", ".json", ".xml", ".yaml", ".yml", ".sql",
    ".c", ".cpp", ".cc", ".h", ".hpp", ".java", ".kt", ".kts", ".cs",
    ".php", ".go", ".rs", ".sh", ".bash", ".bat", ".cmd", ".ps1",
    ".csv", ".log", ".ini", ".toml", ".env",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | IMAGE_EXTENSIONS | {".pdf", ".docx", ".xlsx"}
_speech_lock = threading.Lock()
_speech_thread: threading.Thread | None = None
_speech_session = None
_speech_active = False


def _ensure_storage() -> None:
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    if not INDEX_PATH.exists():
        INDEX_PATH.write_text("{}", encoding="utf-8")


def sanitize_filename(filename: str) -> str:
    """Return a safe display/storage name without path components."""
    raw = str(filename or "upload").replace("\\", "/")
    name = Path(raw).name
    name = re.sub(r"[^A-Za-z0-9._()\[\] -]", "_", name).strip(" .")
    if not name or name in {".", ".."}:
        name = "upload"
    return name[:180]


def _load_index() -> dict[str, dict[str, Any]]:
    _ensure_storage()
    try:
        value = json.loads(INDEX_PATH.read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def _save_index(index: dict[str, dict[str, Any]]) -> None:
    _ensure_storage()
    temporary = INDEX_PATH.with_suffix(".tmp")
    temporary.write_text(json.dumps(index, indent=2), encoding="utf-8")
    temporary.replace(INDEX_PATH)


def _record_path(file_id: str) -> Path:
    path = (FILES_DIR / file_id).resolve()
    if path.parent != FILES_DIR.resolve():
        raise ValueError("Invalid file identifier.")
    return path


def _record(record: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in record.items() if key != "storage_path"}


def file_type(filename: str) -> dict[str, str | bool]:
    extension = Path(filename).suffix.lower()
    if extension in TEXT_EXTENSIONS:
        category = "code" if extension in {".py", ".js", ".ts", ".jsx", ".tsx", ".html", ".css", ".sql", ".c", ".cpp", ".h", ".hpp", ".java", ".kt", ".kts", ".cs", ".php", ".go", ".rs", ".sh", ".bat", ".cmd", ".ps1"} else "text"
        supported = True
    elif extension in {".pdf", ".docx", ".xlsx"}:
        category, supported = "document", True
    elif extension in IMAGE_EXTENSIONS:
        category, supported = "image", True
    else:
        category, supported = "unsupported", False
    return {"extension": extension, "category": category, "supported": supported, "mime": mimetypes.guess_type(filename)[0] or "application/octet-stream"}


def _decode_text(path: Path) -> tuple[str, dict[str, Any]]:
    data = path.read_bytes()
    truncated = len(data) > MAX_TEXT_BYTES
    text = data[:MAX_TEXT_BYTES].decode("utf-8", errors="replace")
    return text, {"truncated": truncated, "bytes": len(data)}


def extract_file(path: Path, extension: str) -> tuple[str, dict[str, Any]]:
    if extension in TEXT_EXTENSIONS:
        return _decode_text(path)
    if extension == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF extraction requires the optional pypdf package.") from exc
        reader = PdfReader(str(path))
        pages = []
        extracted_page_text = []
        for number, page in enumerate(reader.pages, 1):
            page_text = page.extract_text() or ""
            extracted_page_text.append(page_text)
            pages.append(f"[Page {number}]\n{page_text}")
        text = "\n\n".join(pages)
        return text[:MAX_TEXT_BYTES], {"pages": len(reader.pages), "extracted_pages": sum(bool(page.strip()) for page in extracted_page_text), "truncated": len(text) > MAX_TEXT_BYTES}
    if extension == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX extraction requires the optional python-docx package.") from exc
        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return text[:MAX_TEXT_BYTES], {"paragraphs": len(document.paragraphs), "truncated": len(text) > MAX_TEXT_BYTES}
    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX extraction requires the optional openpyxl package.") from exc
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                lines.append(", ".join("" if value is None else str(value) for value in row))
        text = "\n".join(lines)
        return text[:MAX_TEXT_BYTES], {"sheets": workbook.sheetnames, "truncated": len(text) > MAX_TEXT_BYTES}
    if extension in IMAGE_EXTENSIONS:
        try:
            from PIL import Image
            with Image.open(path) as image:
                return "", {"width": image.width, "height": image.height, "format": image.format, "vision_available": False}
        except Exception:
            return "", {"vision_available": False}
    raise ValueError("Unsupported file type.")


def save_uploaded_file(filename: str, content: bytes, *, mime_type: str = "") -> dict[str, Any]:
    if len(content) > MAX_FILE_BYTES:
        raise ValueError(f"File exceeds {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    safe_name = sanitize_filename(filename)
    detected = file_type(safe_name)
    file_id = uuid.uuid4().hex
    storage_path = _record_path(file_id)
    _ensure_storage()
    storage_path.write_bytes(content)
    now = datetime.now(timezone.utc).isoformat()
    record: dict[str, Any] = {
        "id": file_id, "filename": safe_name, "mime_type": mime_type or detected["mime"],
        "type": detected, "size": len(content), "uploaded_at": now,
        "processing_status": "PROCESSING", "extracted_text_available": False,
        "metadata": {}, "storage_path": str(storage_path),
    }
    index = _load_index()
    index[file_id] = record
    _save_index(index)
    try:
        text, metadata = extract_file(storage_path, str(detected["extension"]))
        record.update({"processing_status": "READY", "extracted_text_available": bool(text.strip()) and metadata.get("extracted_pages", 1) > 0, "metadata": metadata})
        if text:
            record["text"] = text
    except Exception as exc:
        record.update({"processing_status": "FAILED", "error": str(exc), "metadata": {}})
    index[file_id] = record
    _save_index(index)
    return _record(record)


def list_files() -> list[dict[str, Any]]:
    index = _load_index()
    return sorted((_record(item) for item in index.values()), key=lambda item: item.get("uploaded_at", ""), reverse=True)


def get_file(file_id: str) -> dict[str, Any]:
    record = _load_index().get(str(file_id))
    if not record:
        raise FileNotFoundError("File context not found.")
    return record


def read_file(file_id: str) -> dict[str, Any]:
    record = get_file(file_id)
    if record.get("processing_status") != "READY":
        raise RuntimeError(record.get("error") or "File is not ready.")
    return {"file": _record(record), "text": record.get("text", ""), "metadata": record.get("metadata", {})}


def search_files(query: str, file_ids: Iterable[str] | None = None) -> list[dict[str, Any]]:
    needle = str(query or "").strip().casefold()
    if not needle:
        return []
    allowed = {str(item) for item in file_ids} if file_ids else None
    matches = []
    for record in _load_index().values():
        if allowed is not None and record.get("id") not in allowed:
            continue
        text = str(record.get("text", ""))
        lines = text.splitlines()
        hits = [{"line": index + 1, "text": line[:500]} for index, line in enumerate(lines) if needle in line.casefold()][:20]
        if hits:
            matches.append({"file": _record(record), "matches": hits, "count": len(hits)})
    return matches


def _safe_edit_path(record: dict[str, Any]) -> Path:
    path = _record_path(str(record["id"]))
    if not path.is_file():
        raise FileNotFoundError("Stored file is missing.")
    return path


def edit_file(file_id: str, find: str, replace: str, *, expected_text: str | None = None) -> dict[str, Any]:
    record = get_file(file_id)
    extension = str(record.get("type", {}).get("extension", ""))
    if extension not in TEXT_EXTENSIONS:
        raise ValueError("Only text and code files support direct editing.")
    if not find:
        raise ValueError("Edit target cannot be empty.")
    path = _safe_edit_path(record)
    original = path.read_text(encoding="utf-8", errors="strict")
    if expected_text is not None and original != expected_text:
        raise ValueError("File changed since it was read; refusing an ambiguous edit.")
    occurrences = original.count(find)
    if occurrences != 1:
        raise ValueError(f"Expected exactly one edit target, found {occurrences}.")
    backup = path.with_name(f"{path.name}.{uuid.uuid4().hex}.bak")
    shutil.copy2(path, backup)
    path.write_text(original.replace(find, replace, 1), encoding="utf-8", newline="")
    index = _load_index()
    record["text"] = path.read_text(encoding="utf-8")
    record["processing_status"] = "READY"
    record["extracted_text_available"] = True
    index[str(file_id)] = record
    _save_index(index)
    updated = read_file(file_id)
    updated["backup"] = str(backup)
    return updated


def remove_file(file_id: str) -> bool:
    index = _load_index()
    record = index.pop(str(file_id), None)
    if not record:
        return False
    _record_path(str(file_id)).unlink(missing_ok=True)
    _save_index(index)
    return True


def summarize_file(file_id: str, question: str = "") -> str:
    from ai.core.service import ai_service
    payload = read_file(file_id)
    text = payload["text"]
    if not text:
        return "This file has no extractable text. Image understanding/OCR is not available through File Intelligence yet."
    instruction = question.strip() or "Create a concise summary with the most important points and page/section references when available."
    context_limit = 120_000
    context = text[:context_limit]
    if len(text) > context_limit:
        context += "\n\n[Context truncated; use file_search for targeted retrieval.]"
    try:
        response = ai_service.generate(
            prompt=f"File: {payload['file']['filename']}\n\n{instruction}\n\nDocument context:\n{context}",
            system_prompt="You are JARVIS File Intelligence. Use only the supplied file context. Do not invent facts. Preserve page and section references.",
            capability="reasoning",
        )
    except Exception as exc:
        print(f"[FILE INTELLIGENCE AI ERROR] {exc}")
        raise RuntimeError("Document analysis is unavailable because no configured AI provider supports this operation.") from exc
    if not response.success:
        print(f"[FILE INTELLIGENCE AI ERROR] {response.error}")
        raise RuntimeError("Document analysis is unavailable because no configured AI provider supports this operation.")
    return response.text


def _speech_chunks(text: str, limit: int = 1200) -> list[str]:
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    chunks: list[str] = []
    for paragraph in paragraphs:
        while len(paragraph) > limit:
            boundary = paragraph.rfind(" ", 0, limit)
            boundary = boundary if boundary > 0 else limit
            chunks.append(paragraph[:boundary].strip())
            paragraph = paragraph[boundary:].strip()
        if paragraph:
            chunks.append(paragraph)
    return chunks


def _speak_worker(text: str, session) -> None:
    from voice.manager import speak, start_speech_session

    try:
        for chunk in _speech_chunks(text):
            if session.cancel_event.is_set():
                break
            speak(chunk, wait=True, session=session, notify_remote=False, record_conversation=False)
    finally:
        global _speech_active
        with _speech_lock:
            if _speech_session is session:
                _speech_active = False


def speak_file(file_id: str, text_override: str = "") -> dict[str, Any]:
    """Speak the active extracted content through JARVIS's existing TTS."""
    global _speech_thread, _speech_session, _speech_active
    payload = read_file(file_id)
    text = str(text_override or payload.get("text", "")).strip()
    if not text:
        raise RuntimeError("This file has no extractable text to speak.")
    with _speech_lock:
        if _speech_active:
            return {"started": False, "busy": True}
        from voice.manager import start_speech_session
        _speech_session = start_speech_session()
        _speech_active = True
        _speech_thread = threading.Thread(target=_speak_worker, args=(text, _speech_session), name="JARVIS-File-Speech", daemon=True)
        _speech_thread.start()
    return {"started": True, "busy": False}


def file_speech_status() -> dict[str, bool]:
    with _speech_lock:
        return {"speaking": bool(_speech_active)}


def stop_file_speech() -> dict[str, bool]:
    global _speech_active
    with _speech_lock:
        session = _speech_session
        was_active = bool(_speech_active)
        _speech_active = False
    if session is not None:
        session.cancel_event.set()
        try:
            from voice.state import current_session
            if current_session() is session:
                from voice.player import stop as stop_audio
                stop_audio()
        except Exception as exc:
            print(f"[FILE INTELLIGENCE SPEECH STOP] {exc}")
    return {"stopped": was_active}


def _active_file_id() -> str:
    files = list_files()
    if not files:
        raise FileNotFoundError("No uploaded file is active.")
    return str(files[0]["id"])


def file_intelligence_action(data: dict[str, Any] | None = None) -> Any:
    data = data or {}
    action = data.get("action", "")
    if action == "file_list": return list_files()
    if action in {"file_intelligence_info", "file_read"}: return read_file(str(data.get("file_id"))) if action == "file_read" else _record(get_file(str(data.get("file_id"))))
    if action == "file_search": return search_files(str(data.get("query", "")), data.get("file_ids"))
    if action in {"file_summarize", "file_question", "file_analyze"}: return summarize_file(str(data.get("file_id")), str(data.get("question", "")))
    if action == "file_intelligence_request": return summarize_file(_active_file_id(), str(data.get("request", "")))
    if action == "file_question_request": return summarize_file(_active_file_id(), str(data.get("question", "")))
    if action == "file_search_request": return search_files(str(data.get("query", "")), [_active_file_id()])
    if action == "file_edit": return edit_file(str(data.get("file_id")), str(data.get("find", "")), str(data.get("replace", "")), expected_text=data.get("expected_text"))
    if action == "file_save": return edit_file(str(data.get("file_id")), str(data.get("find", "")), str(data.get("replace", "")), expected_text=data.get("expected_text"))
    if action == "file_preview": return read_file(str(data.get("file_id")))
    if action == "file_speak": return speak_file(str(data.get("file_id")))
    if action == "file_stop_speech": return stop_file_speech()
    if action == "file_open_folder":
        path = _safe_edit_path(get_file(str(data.get("file_id"))))
        if os.name == "nt":
            subprocess.Popen(["explorer.exe", f"/select,{path}"])
        return str(path)
    if action == "file_remove": return remove_file(str(data.get("file_id")))
    raise ValueError("Unknown File Intelligence action.")


for _action in ("file_list", "file_intelligence_info", "file_read", "file_search", "file_summarize", "file_question", "file_analyze", "file_edit", "file_save", "file_preview", "file_speak", "file_open_folder", "file_intelligence_request", "file_question_request", "file_search_request"):
    register(_action, file_intelligence_action, category="files")
